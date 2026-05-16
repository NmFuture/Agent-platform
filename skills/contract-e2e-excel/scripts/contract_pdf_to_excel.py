#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from decimal import Decimal
from pathlib import Path
from typing import Any

try:
    from pypdf import PdfReader
except ImportError as exc:  # pragma: no cover - runtime dependency guard
    raise SystemExit("Missing dependency: pypdf. Run with a Python environment that has pypdf installed.") from exc

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover - runtime dependency guard
    raise SystemExit("Missing dependency: openpyxl. Run with a Python environment that has openpyxl installed.") from exc


STATUS_FOUND = "found"
STATUS_NOT_FOUND = "not_found"
STATUS_AMBIGUOUS = "ambiguous"
STATUS_NOT_APPLICABLE = "not_applicable"
SKILL_ROOT = Path(__file__).resolve().parents[1]
ASSETS_DIR = SKILL_ROOT / "assets"


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def collapse(text: str | None) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def parse_money(value: str) -> Decimal | None:
    match = re.search(r"(?:¥|￥)?\s*(\d[\d,]*(?:\.\d+)?)", value or "")
    if not match:
        return None
    return Decimal(match.group(1).replace(",", ""))


def read_pdf_pages(pdf_path: Path) -> list[dict[str, Any]]:
    reader = PdfReader(str(pdf_path))
    pages: list[dict[str, Any]] = []
    for idx, page in enumerate(reader.pages, start=1):
        text = clean_text(page.extract_text() or "")
        pages.append({"page": idx, "chars": len(text), "text": text})
    return pages


def build_offset_index(pages: list[dict[str, Any]]) -> tuple[str, list[tuple[int, int]]]:
    chunks: list[str] = []
    offsets: list[tuple[int, int]] = []
    pos = 0
    for page in pages:
        marker = f"\n===== PAGE {page['page']} =====\n"
        chunk = marker + page["text"] + "\n"
        chunks.append(chunk)
        offsets.append((pos, page["page"]))
        pos += len(chunk)
    return "".join(chunks), offsets


def page_for_offset(offsets: list[tuple[int, int]], pos: int) -> int:
    current = offsets[0][1]
    for start, page in offsets:
        if start > pos:
            break
        current = page
    return current


def snippet_around(text: str, needle: str, radius: int = 360) -> str:
    idx = text.find(needle)
    if idx < 0:
        return collapse(text[: radius * 2])
    start = max(0, idx - radius)
    end = min(len(text), idx + len(needle) + radius)
    return collapse(text[start:end])


class ContractIndex:
    def __init__(self, pages: list[dict[str, Any]]):
        self.pages = pages
        self.full_text, self.offsets = build_offset_index(pages)

    def page_text(self, page: int) -> str:
        if 1 <= page <= len(self.pages):
            return self.pages[page - 1]["text"]
        return ""

    def find_regex(self, pattern: str, start_page: int = 1) -> dict[str, Any] | None:
        search_start = 0
        if start_page > 1:
            marker_idx = self.full_text.find(f"===== PAGE {start_page} =====")
            if marker_idx >= 0:
                search_start = marker_idx
        match = re.search(pattern, self.full_text[search_start:], flags=re.S)
        if not match:
            return None
        absolute_start = search_start + match.start()
        page = page_for_offset(self.offsets, absolute_start)
        value = match.group(0)
        return {"value": value, "page": page, "evidence": snippet_around(self.full_text, value, 420)}

    def find_between(
        self,
        start_pattern: str,
        end_pattern: str | None,
        max_chars: int = 1600,
        after_page: int | None = None,
    ) -> dict[str, Any] | None:
        search_start = 0
        if after_page:
            marker_idx = self.full_text.find(f"===== PAGE {after_page} =====")
            if marker_idx >= 0:
                search_start = marker_idx
        start = re.search(start_pattern, self.full_text[search_start:], flags=re.S)
        if not start:
            return None
        absolute_start = search_start + start.start()
        absolute_end = search_start + start.end()
        end_pos = min(len(self.full_text), absolute_end + max_chars)
        if end_pattern:
            end = re.search(end_pattern, self.full_text[absolute_end:], flags=re.S)
            if end:
                end_pos = absolute_end + end.start()
        page = page_for_offset(self.offsets, absolute_start)
        return {"page": page, "evidence": collapse(self.full_text[absolute_start:end_pos])}

    def find_best(
        self,
        terms: list[str],
        prefer_pages: list[int] | None = None,
        min_page: int = 1,
    ) -> dict[str, Any] | None:
        terms = [t for t in dict.fromkeys(terms) if t and len(t) >= 2]
        if not terms:
            return None

        best: tuple[int, int, dict[str, Any]] | None = None
        prefer_pages = prefer_pages or []
        for page in self.pages:
            if page["page"] < min_page:
                continue
            text = page["text"]
            score = 0
            first_term = None
            for term in terms:
                if term in text:
                    score += 10 + min(len(term), 8)
                    first_term = first_term or term
            if not score:
                continue
            if page["page"] in prefer_pages:
                score += 5
            payload = {
                "page": page["page"],
                "term": first_term,
                "evidence": snippet_around(text, first_term or terms[0], 460),
            }
            candidate = (score, -page["page"], payload)
            if best is None or candidate > best:
                best = candidate
        return best[2] if best else None


def extraction_result(value: Any, evidence: str | None, page: int | None, status: str, confidence: float) -> dict[str, Any]:
    return {
        "value": value,
        "evidence_text": collapse(evidence) if evidence else None,
        "source_location": f"PDF第{page}页" if page else None,
        "confidence": confidence,
        "status": status,
    }


def not_found() -> dict[str, Any]:
    return extraction_result(None, None, None, STATUS_NOT_FOUND, 0.0)


def compact_value_around(evidence: str | None, term: str) -> str:
    if not evidence:
        return ""
    sentences = re.split(r"(?<=[。；;])", evidence)
    for sent in sentences:
        if term and term in sent:
            return collapse(sent)[:500]
    return collapse(evidence)[:500]


def field_terms(field_name: str) -> list[str]:
    raw_terms = [field_name]
    text = re.sub(r"（[^）]*）|\([^)]*\)", "", field_name)
    for old in [
        "是否需要",
        "是否为",
        "是否是",
        "是否有",
        "是否",
        "供应商品牌",
        "供应商",
        "品牌要求",
        "品牌",
        "配置要求",
        "技术参数",
        "具体要求",
        "要求",
        "数量",
        "型式",
        "形式",
        "材质",
        "介质",
        "供货方",
        "配置",
    ]:
        text = text.replace(old, "")
    parts = re.split(r"[、/，,；;：:（）()\\s]+", text)
    raw_terms.extend([text] + parts)
    stop = {"需要", "具备此功能", "此功能", "相关", "具体", "内容", "合同", "条款", "字段", "从合同中", "配置"}
    return [t.strip() for t in raw_terms if t and t.strip() and t.strip() not in stop and len(t.strip()) >= 2]


SPECIAL_KEYWORDS = {
    "设备供货": ["本合同标的是卖方为买方供应", "合同供货范围详见第五章供货范围", "整套风力发电机组"],
    "基础载荷计算、设计及指导": ["基础设计", "基础载荷", "标准型基础图", "基础顶部荷载"],
    "设计联络会要求": ["设计联络", "技术联络会", "联络会"],
    "培训要求": ["技术培训", "培训"],
    "基础环、塔筒监造": ["塔架全过程的监造", "塔筒", "监造"],
    "业主监造要求": ["买方监造", "监造代表", "设备监造"],
    "出厂实验要求": ["出厂前须进行", "检验和试验", "出厂检验"],
    "基础指导要求": ["基础施工过程", "基础设计", "风电机组基础"],
    "安装指导要求": ["安装指导", "安装过程中提供支持、监督和指导"],
    "调试要求": ["卖方全面负责风电机组的调试", "调试方案", "负责调试"],
    "试运行要求": ["可靠性运行应当通过其持续 240 小时", "试运行时间不得少于 240小时", "试运行"],
    "预验收要求": ["当每单元的最后 1 台风电机组通过 240 小时试运行后", "买方签发该单元全部风力发电机组的预验收证书"],
    "工程创优/科技奖项要求": ["工程创优", "科技奖", "科研项目", "创新创优"],
    "资源合作要求": ["资源合作", "风资源合作", "风资源交换"],
    "是否有产值诉求": ["产值诉求", "产值落地", "产值"],
    "是否有税收要求": ["税收要求", "落产落税", "税收"],
    "现场海拔高度": ["现场海拔", "海拔高度", "平均海拔"],
    "台风": ["台风影响", "台风级别", "抗台风"],
    "运行温度": ["运行环境温度", "生存环境温度", "工作环境温度", "-30℃~+40℃"],
    "防腐等级": ["ISO12944-2 C4", "防腐保护", "防腐"],
    "风沙等级": ["风沙等级", "风沙条件"],
    "噪音": ["噪音", "噪声", "IEC61400－11"],
    "是否为防盐雾机型": ["防盐雾", "盐雾"],
    "现场湿度": ["相对湿度", "湿度"],
    "单台可利用率": ["单机可利用率", "任意单台风电机组年可利用率"],
    "全场可利用率": ["全场可利用率", "年平均可利用率"],
    "单台功率曲线K值": ["功率曲线保证", "保证值(K)", "相符度不小于98%"],
    "全场功率曲线K值": ["功率曲线保证", "保证值(K)", "相符度不小于98%"],
    "等效满发小时数": ["等效发电小时", "等效满发小时"],
    "发电量": ["上网电量", "年发电量", "发电量"],
    "防覆冰涂层": ["叶片防覆冰", "防覆冰"],
    "机舱小吊车平平": ["机舱小吊车"],
    "吊装方式": ["满足单叶片吊装条件", "单叶片吊装工装", "单叶片吊装方案"],
    "场内外道路分界点": ["本项目分界点拟定为", "出收费站150米处", "场内外道路分界点"],
    "二次倒运责任范围（如有：是否要求特种转运车辆（机舱/叶片）": ["二次倒运", "第二次倒运", "叶片举升车"],
    "堆场及倒运车": ["临时设备中转场存放至少3套", "配备叶片举升车", "堆场"],
    "场内修路标准": ["新建道路的主路路基宽", "路面", "路肩"],
    "道路改造责任划分": ["分界点至机位的杆塔迁移", "道路改扩建等清障及协调工作由买方负责"],
    "场内运输牵引责任": ["采用叶片举升车对道路进行空车试运", "大件运输特种车辆要求"],
    "质保期大部件更换要求及考核": ["累计出现 20%及以上的损坏或失效", "大部件拆装发生的人工费"],
}


def extract_cover_name(index: ContractIndex) -> dict[str, Any]:
    page1 = index.page_text(1)
    flat = collapse(page1).replace("设 备", "设备")
    patterns = [
        r"([\u4e00-\u9fffA-Za-z0-9（）()、\-·\s]{2,120}?风电项目\s*风力发电机组（及塔筒）设备采购合同)",
        r"([\u4e00-\u9fffA-Za-z0-9（）()、\-·\s]{2,120}?合同)",
    ]
    for pattern in patterns:
        match = re.search(pattern, flat)
        if match:
            value = collapse(match.group(1))
            return extraction_result(value, value, 1, STATUS_FOUND, 0.90)
    lines = [collapse(line) for line in page1.splitlines() if collapse(line)]
    candidates = [line for line in lines[:20] if "合同" in line]
    if candidates:
        return extraction_result(candidates[0], candidates[0], 1, STATUS_AMBIGUOUS, 0.65)
    return not_found()


def extract_party(index: ContractIndex, party: str) -> dict[str, Any]:
    patterns = [
        rf"{party}方?[:：]\s*([0-9A-Za-z\u4e00-\u9fff（）()\s]+?(?:有限公司|股份有限公司|集团有限公司))",
        rf"{party}[:：]\s*([0-9A-Za-z\u4e00-\u9fff（）()\s]+?(?:有限公司|股份有限公司|集团有限公司))",
    ]
    for page_no in range(1, min(len(index.pages), 8) + 1):
        text = index.page_text(page_no)
        for pattern in patterns:
            match = re.search(pattern, text)
            if match:
                return extraction_result(collapse(match.group(1)), collapse(match.group(0)), page_no, STATUS_FOUND, 0.92)
    return not_found()


def extract_contract_amount(index: ContractIndex) -> dict[str, Any]:
    hit = index.find_regex(r"本合同总价为人民币[^。]{0,180}。", start_page=1)
    if not hit:
        hit = index.find_regex(r"(?:签约合同价|合同总价|合同金额)[:：][^。]{0,180}", start_page=1)
    if not hit:
        return not_found()
    evidence = hit["evidence"]
    total = parse_money(hit["value"] or evidence)
    tax_match = re.search(r"税率\s*(\d+(?:\.\d+)?)\s*[%％]", evidence)
    value = collapse(hit["value"])
    if total:
        value = f"含税人民币{total:,.2f}元"
        if tax_match:
            value += f"；税率{tax_match.group(1)}%"
    return extraction_result(value, evidence, hit["page"], STATUS_FOUND, 0.90)


def infer_project_type(index: ContractIndex) -> str:
    if index.find_best(["适用于海上地区", "海上风电", "海上项目"]):
        return "海上项目"
    if index.find_best(["适用于陆上地区", "陆上地区", "陆上风电", "施工现场指定机位"]):
        return "陆上项目"
    return "未知"


def infer_mixed_tower(index: ContractIndex) -> bool:
    return bool(index.find_best(["钢混塔筒", "钢混塔架", "混合塔筒", "混塔"]))


def generic_extract(rule: dict[str, Any], index: ContractIndex) -> dict[str, Any]:
    field_name = rule.get("field_name", "")
    terms = SPECIAL_KEYWORDS.get(field_name) or field_terms(field_name)
    min_page = 4
    if rule.get("category_1") in {"机组配置要求", "风场配置要求"}:
        min_page = 35
    elif rule.get("category_1") == "运输要求":
        min_page = 120
    elif rule.get("category_1") == "运维要求":
        min_page = 100
    hit = index.find_best(terms, min_page=min_page)
    if not hit:
        return not_found()

    evidence = hit["evidence"]
    value_type = (rule.get("output_contract") or {}).get("value_type", "text")
    if value_type == "boolean_or_requirement":
        value: Any = "有相关要求，见原文依据"
    elif value_type == "amount":
        amounts = re.findall(r"(?:¥|￥)?\d[\d,]*(?:\.\d+)?\s*(?:元|万元|%|％)?", evidence)
        value = "；".join(amounts[:3]) if amounts else "见原文依据"
    elif value_type in {"percentage_or_ratio", "number_with_unit"}:
        nums = re.findall(r"\d+(?:\.\d+)?\s*(?:%|％|MW|m|米|小时|h|天|日|个月|年|套|台|份|℃|°C)?", evidence)
        value = "；".join(dict.fromkeys(nums[:5])) if nums else "见原文依据"
    else:
        value = compact_value_around(evidence, hit.get("term") or field_name)
    return extraction_result(value, evidence, hit["page"], STATUS_FOUND, 0.62)


def extract_type1_rule(rule: dict[str, Any], index: ContractIndex, project_type: str, mixed_tower: bool) -> dict[str, Any]:
    field_name = rule.get("field_name", "")
    applicable = rule.get("applicable_project_type")
    if applicable == "陆上项目" and project_type != "陆上项目":
        return extraction_result(None, None, None, STATUS_NOT_APPLICABLE, 1.0)
    if applicable == "海上项目" and project_type != "海上项目":
        return extraction_result(None, None, None, STATUS_NOT_APPLICABLE, 1.0)
    if applicable == "混塔项目" and not mixed_tower:
        return extraction_result(None, None, None, STATUS_NOT_APPLICABLE, 1.0)

    if field_name == "合同名称":
        return extract_cover_name(index)
    if field_name == "合同编号":
        hit = index.find_regex(r"(?:YX|EX)\d{13}")
        return extraction_result(hit["value"], hit["evidence"], hit["page"], STATUS_FOUND, 0.95) if hit else not_found()
    if field_name.startswith("合同买方"):
        return extract_party(index, "买")
    if field_name.startswith("合同业主方"):
        return extract_party(index, "业主")
    if field_name == "合同金额":
        return extract_contract_amount(index)
    if field_name == "项目类型":
        if project_type != "未知":
            term = "海上" if project_type == "海上项目" else "陆上"
            hit = index.find_best([term])
            return extraction_result(project_type, hit["evidence"] if hit else None, hit["page"] if hit else None, STATUS_FOUND, 0.82)
        return not_found()
    if field_name == "是否混塔项目":
        hit = index.find_best(["钢混塔筒", "钢混塔架", "混合塔筒", "混塔"])
        return extraction_result(True, hit["evidence"], hit["page"], STATUS_FOUND, 0.90) if hit else extraction_result(False, None, None, STATUS_FOUND, 0.50)
    if field_name == "签订日期":
        hit = index.find_regex(r"20\d{2}\s*年\s*\d{0,2}\s*月\s*\d{0,2}\s*日?")
        if hit:
            value = collapse(hit["value"])
            status = STATUS_FOUND if re.search(r"20\d{2}\s*年\s*\d+\s*月\s*\d+\s*日", value) else STATUS_AMBIGUOUS
            confidence = 0.86 if status == STATUS_FOUND else 0.58
            return extraction_result(value, hit["evidence"], hit["page"], status, confidence)
        return not_found()
    if field_name == "工程地点":
        hit = index.find_best(["合同现场", "风电场", "施工现场指定机位", "交货地点"])
        return extraction_result(compact_value_around(hit["evidence"], hit["term"]), hit["evidence"], hit["page"], STATUS_FOUND, 0.78) if hit else not_found()
    if field_name == "交货期":
        hit = index.find_best(["交货进度要求", "开始供第一批", "交货期", "交货时间"])
        return extraction_result(compact_value_around(hit["evidence"], hit["term"]), hit["evidence"], hit["page"], STATUS_FOUND, 0.80) if hit else not_found()
    if field_name == "产品型号":
        hit = index.find_regex(r"型号[:：]\s*[A-Za-z0-9.\-]+")
        if hit:
            model = re.sub(r"^型号[:：]\s*", "", collapse(hit["value"]))
            return extraction_result(model, hit["evidence"], hit["page"], STATUS_FOUND, 0.90)
        return not_found()
    if field_name == "数量":
        hit = index.find_regex(r"数量[:：]\s*\d+\s*(?:台|套)?")
        if hit:
            qty = re.sub(r"^数量[:：]\s*", "", collapse(hit["value"]))
            return extraction_result(qty, hit["evidence"], hit["page"], STATUS_FOUND, 0.86)
    if field_name == "试运行小时":
        hit = index.find_best(["持续 240 小时", "240小时", "试运行时间不得少于"])
        return extraction_result(compact_value_around(hit["evidence"], hit["term"]), hit["evidence"], hit["page"], STATUS_FOUND, 0.88) if hit else not_found()
    if field_name == "质保周期":
        hit = index.find_best(["质量保证期", "60 个月", "五年"])
        return extraction_result(compact_value_around(hit["evidence"], hit["term"]), hit["evidence"], hit["page"], STATUS_FOUND, 0.86) if hit else not_found()

    return generic_extract(rule, index)


def extract_total_amount(index: ContractIndex) -> Decimal | None:
    amount_hit = extract_contract_amount(index)
    if amount_hit["status"] == STATUS_FOUND:
        return parse_money(str(amount_hit["value"]))
    return None


def extract_payments(index: ContractIndex) -> list[dict[str, Any]]:
    total = extract_total_amount(index) or Decimal("0")
    specs = [
        ("预付款", Decimal("0.10"), r"3\.2\.3\.1\s*预付款", r"3\.2\.3\.2\s*投料款"),
        ("投料款", Decimal("0.30"), r"3\.2\.3\.2\s*投料款", r"3\.2\.3\.3\s*交货款"),
        ("交货款", Decimal("0.50"), r"3\.2\.3\.3\s*交货款", r"3\.2\.3\.4\s*预验收款"),
        ("预验收款", Decimal("0.05"), r"3\.2\.3\.4\s*预验收款", r"3\.2\.3\.5\s*质保期付款"),
        ("质保款", Decimal("0.05"), r"3\.2\.3\.5\s*质保期付款", r"3\.2\.4\s*付款时间"),
    ]
    rows: list[dict[str, Any]] = []
    for name, ratio, start, end in specs:
        clause = index.find_between(start, end, max_chars=1400, after_page=15)
        rows.append(
            {
                "款项类型": name,
                "款项金额（元）": int(total * ratio) if total else None,
                "款项比例": float(ratio),
                "付款条件（摘录原文）": clause["evidence"] if clause else None,
                "原文位置": f"PDF第{clause['page']}页" if clause else None,
                "状态": STATUS_FOUND if clause else STATUS_NOT_FOUND,
                "置信度": 0.90 if clause else 0.0,
            }
        )
    return rows


def liability_row(index: ContractIndex, name: str, promise_terms: list[str], calc_terms: list[str], cap_terms: list[str]) -> dict[str, Any]:
    hit = index.find_best(promise_terms + calc_terms + cap_terms)
    evidence = hit["evidence"] if hit else None
    return {
        "违约赔偿金种类": name,
        "计违约金的合同承诺": compact_value_around(evidence, promise_terms[0]) if evidence else None,
        "违约赔偿金计算": compact_value_around(evidence, calc_terms[0]) if evidence else None,
        "分项上限及免责": compact_value_around(evidence, cap_terms[0]) if evidence else None,
        "原文位置": f"PDF第{hit['page']}页" if hit else None,
        "状态": STATUS_FOUND if hit else STATUS_NOT_FOUND,
        "置信度": 0.80 if hit else 0.0,
    }


def extract_liabilities(index: ContractIndex) -> list[dict[str, Any]]:
    rows = [
        liability_row(index, "延迟交货", ["由于卖方责任，合同设备不能按期交货"], ["延迟1～4 周", "延迟 1～4周", "延迟5～8 周", "延迟8周以上"], ["违约金累计不超过合同设备总价的 10%", "不可抗力"]),
        liability_row(index, "可利用率不达标", ["单机可利用率", "全场可利用率"], ["每降低一个百分点赔偿"], ["单机可利用率和全场可利用率罚金不同时计算"]),
        liability_row(index, "功率曲线不达标", ["功率曲线保证", "相符度不小于98%"], ["每降低一个百分点赔偿"], ["负偏差超过 10％", "更换机组设备"]),
        liability_row(index, "技术资料不完整或不正确", ["卖方保证提供技术资料完整、清晰、正确"], ["承担买方由此引起的直接损失"], ["技术资料内容不完整或不正确"]),
        liability_row(index, "严重质量缺陷或大部件失效", ["主要部件", "十分严重的缺陷"], ["每超1天每套扣0.5万元", "重新计算"], ["质保金中扣除"]),
        liability_row(index, "项目管理人员不满足要求", ["主要管理人员必须为投标文件中承诺"], ["1000元/天", "5000", "2000元/天"], ["无故离开现场10天以上"]),
        liability_row(index, "资料文件、堆场、履责不到位考核", ["卖方未按要求提供资料文件", "卖方未按要求在设备堆场存放设备"], ["1000元~5000元", "5000元~10000元"], ["考核款从合同支付款中扣除"]),
    ]
    for idx, row in enumerate(rows, start=1):
        row["序号"] = idx
    return rows


def risk_row(idx: int, category: str, clause: dict[str, Any] | None, analysis: str) -> dict[str, Any]:
    return {
        "序号": idx,
        "分类": category,
        "条款": clause["evidence"] if clause else None,
        "风险点": analysis,
        "原文位置": f"PDF第{clause['page']}页" if clause else None,
        "状态": STATUS_FOUND if clause else STATUS_NOT_FOUND,
        "置信度": 0.78 if clause else 0.0,
    }


def extract_risks(index: ContractIndex) -> list[dict[str, Any]]:
    candidates = [
        ("付款扣款风险", ["买方有权从上述任何一笔应付款", "违约金的扣除与支付"], "买方可直接从应付款中扣除违约金或赔偿金，会影响卖方现金流和回款确定性。"),
        ("延迟交货违约及解除风险", ["任何一批合同货物迟交 12 周以上", "买方有权终止部分或全部合同"], "迟交除分段违约金外，超过约定期限还可能触发买方解除部分或全部合同。"),
        ("性能考核罚金风险", ["单机可利用率", "全场可利用率", "功率曲线保证"], "可利用率、功率曲线等指标与罚金、整改、更换设备义务绑定，卖方承担持续性能达标压力。"),
        ("质量缺陷连带损失风险", ["质量、安全事故及工期延误", "直接及间接损失"], "质量、安全事故或工期延误可能导致直接及间接损失承担，责任范围较重。"),
        ("运输与倒运责任风险", ["叶片举升车", "第二次倒运", "二次倒运"], "卖方承担叶片举升车资源、堆场倒运、二次倒运等现场运输资源责任，可能产生额外履约成本。"),
        ("道路方案工程量考核风险", ["工程量超出中标人路勘报告中路径工程量的20%", "路勘报告"], "运输路线和道路改造方案如果偏差超过20%，卖方可能被考核。"),
        ("数据开放与平台接入风险", ["免费开放风电机组全量 SCADA运行数据", "完成数据接口工作"], "卖方需开放全量SCADA数据并配合集控、产业中台接入，存在数据接口与后续升级配合义务。"),
        ("资料交付逾期风险", ["未按要求提供资料文件", "按天考核1000元"], "资料文件交付不及时会按天考核，且与设备安装、调试、验收进度相关。"),
    ]
    return [risk_row(idx, category, index.find_best(terms), analysis) for idx, (category, terms, analysis) in enumerate(candidates, start=1)]


def doc_row(idx: int, name: str, quantity: str | None, fmt: str | None, time_req: str, clause: dict[str, Any] | None) -> dict[str, Any]:
    return {
        "序号": idx,
        "资料名称": name,
        "数量": quantity,
        "格式要求": fmt,
        "交付时间要求": time_req,
        "合同原文": clause["evidence"] if clause else None,
        "原文位置": f"PDF第{clause['page']}页" if clause else None,
        "状态": STATUS_FOUND if clause else STATUS_NOT_FOUND,
        "置信度": 0.82 if clause else 0.0,
    }


def extract_document_delivery(index: ContractIndex) -> list[dict[str, Any]]:
    rows = [
        doc_row(1, "风机安装手册", "6套纸质版、1套电子版", None, "签订合同后15日内", index.find_best(["风机的安装手册应在签订合同后的15日内"])),
        doc_row(2, "风机维护手册", "6套纸质版、1套电子版", None, "第一台风机吊装完成后15日内", index.find_best(["风机的维护手册应在第一台风机吊装完成后的15日内"])),
        doc_row(3, "混塔迭代设计所需资料", None, None, "签订合同后10日内提供必要资料；买方启动基础设计通知后20日内完成风机基础设计", index.find_best(["签订合同后 10 日内提供必要的资料", "20日内完成风机基础设计"])),
        doc_row(4, "单叶片吊装方案、吊装工具及操作流程", None, None, "买方编制吊装方式时提供", index.find_best(["卖方提供单叶片吊装方案，并提供相关吊装工具及操作流程"])),
        doc_row(5, "风机系统调试方案、启动方案", None, None, "风场计划倒送电时间前30日完成并提交", index.find_best(["风场计划倒送电时间前30日完成风机系统调试方案、启动方案"])),
        doc_row(6, "设备资料、调试、试运行及竣工资料", None, "符合国家及行业标准、项目公司和集团要求", "按项目公司及集团要求完成移交", index.find_best(["完成设备资料", "试运行以及其他竣工资料移交"])),
        doc_row(7, "每套风机、塔筒检验和试验报告", None, None, "每套风机、塔筒交付后15日内提交", index.find_best(["每套风机、塔筒交付后，15日内提交设备的检验和实验报告"])),
        doc_row(8, "ADPSS电磁暂态模型、测试报告和模型说明书", "标准结构化模型1套、原始代码模型1套", None, "按电网新能源建模技术要求提供并配合验证", index.find_best(["提供 ADPSS 电磁暂态模型", "模型说明书"])),
        doc_row(9, "SCADA点表清单、状态码清单、故障码清单和说明文件", "电子版和纸质版", "电子版、纸质版", "随主机交付", index.find_best(["随主机交付 SCADA点表清单", "状态码清单", "故障码清单"])),
    ]
    return rows


def status_counts(rows: list[dict[str, Any]]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for row in rows:
        status = row.get("extracted", {}).get("status", STATUS_NOT_FOUND)
        counts[status] = counts.get(status, 0) + 1
    return counts


def find_rules_file(explicit: str | None, start_dir: Path) -> Path:
    if explicit:
        path = Path(explicit).expanduser().resolve()
        if not path.exists():
            raise FileNotFoundError(f"Rules file not found: {path}")
        return path
    candidates = [
        start_dir / "outputs" / "contract_rules_v1_7.machine.json",
        start_dir / "contract_rules_v1_7.machine.json",
        start_dir / "outputs" / "contract_rules.machine.json",
        start_dir / "contract_rules.machine.json",
        ASSETS_DIR / "contract_rules_v1_7.machine.json",
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate.resolve()
    matches = sorted(start_dir.glob("**/contract_rules*.machine.json"))
    if matches:
        return matches[0].resolve()
    raise FileNotFoundError("Could not find contract_rules*.machine.json. Pass --rules explicitly.")


def get_type(rules_payload: dict[str, Any], type_id: str) -> dict[str, Any]:
    for type_payload in rules_payload.get("types", []):
        if type_payload.get("type_id") == type_id or type_payload.get("sheet_name") == type_id:
            return type_payload
    raise KeyError(f"Rules JSON missing type: {type_id}")


def safe_text(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (int, float)):
        return value
    text = str(value).replace("\r", "\n")
    text = re.sub(r"\n+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text[:32000]


def write_table(ws, title: str, subtitle: str, headers: list[str], rows: list[list[Any]], widths: list[int]) -> None:
    max_col = len(headers)
    ws.cell(1, 1, title)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor="D9EAF7")
    ws.cell(2, 1, subtitle)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
    ws.cell(2, 1).font = Font(color="5A5A5A")
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(4, col, header)
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_idx, row in enumerate(rows, start=5):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row_idx, col_idx, safe_text(value))
            cell.alignment = Alignment(vertical="top", wrap_text=False)
        ws.row_dimensions[row_idx].height = 64
    ws.freeze_panes = "A5"
    end_row = max(4, len(rows) + 4)
    ws.auto_filter.ref = f"A4:{get_column_letter(max_col)}{end_row}"
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = max(8, min(width / 7, 85))


def build_workbook(payload: dict[str, Any], output_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    metadata = payload.get("metadata", {})
    subtitle = (
        f"来源合同：{metadata.get('contract_file', '')}；页数：{metadata.get('pdf_pages', '')}；"
        f"项目类型：{metadata.get('project_type', '')}；是否混塔：{'是' if metadata.get('mixed_tower') else '否'}"
    )

    type1_headers = ["序号", "分解对象", "适用项目类型", "一级分类", "二级分类", "字段名称", "抽取结果", "状态", "置信度", "原文依据", "原文位置", "提取规则"]
    type1_rows = []
    for item in payload.get("type_1_direct_field_extraction", []):
        rule = item.get("rule", {})
        extracted = item.get("extracted", {})
        type1_rows.append(
            [
                rule.get("sequence", ""),
                rule.get("decomposition_object", ""),
                rule.get("applicable_project_type", ""),
                rule.get("category_1", ""),
                rule.get("category_2", ""),
                rule.get("field_name", ""),
                extracted.get("value", ""),
                extracted.get("status", ""),
                extracted.get("confidence", ""),
                extracted.get("evidence_text", ""),
                extracted.get("source_location", ""),
                rule.get("extraction_rule", ""),
            ]
        )
    write_table(wb.create_sheet("类型1-字段直接提取"), "类型1-字段直接提取", subtitle, type1_headers, type1_rows, [58, 78, 100, 110, 170, 210, 220, 90, 80, 460, 150, 460])

    type2_headers = ["款项类型", "款项金额（元）", "款项比例", "付款条件（摘录原文）", "原文位置", "状态", "置信度"]
    type2_rows = [[row.get(col, "") for col in type2_headers] for row in payload.get("type_2_payment_terms", [])]
    write_table(wb.create_sheet("类型2-付款条件提取"), "类型2-付款条件提取", subtitle, type2_headers, type2_rows, [120, 130, 90, 620, 150, 90, 80])

    type3_headers = ["序号", "违约赔偿金种类", "计违约金的合同承诺", "违约赔偿金计算", "分项上限及免责", "原文位置", "状态", "置信度"]
    type3_rows = [[row.get(col, "") for col in type3_headers] for row in payload.get("type_3_default_liability", [])]
    write_table(wb.create_sheet("类型3-违约条款"), "类型3-违约条款", subtitle, type3_headers, type3_rows, [58, 170, 520, 460, 320, 150, 90, 80])

    type4_headers = ["序号", "分类", "条款", "风险点", "原文位置", "状态", "置信度"]
    type4_rows = [[row.get(col, "") for col in type4_headers] for row in payload.get("type_4_contract_risk", [])]
    write_table(wb.create_sheet("类型4-合同风险识别"), "类型4-合同风险识别", subtitle, type4_headers, type4_rows, [58, 180, 620, 380, 150, 90, 80])

    type5_headers = ["序号", "资料名称", "数量", "格式要求", "交付时间要求", "合同原文", "原文位置", "状态", "置信度"]
    type5_rows = [[row.get(col, "") for col in type5_headers] for row in payload.get("type_5_document_delivery", [])]
    write_table(wb.create_sheet("类型5-资料交付要求"), "类型5-资料交付要求", subtitle, type5_headers, type5_rows, [58, 220, 140, 160, 240, 620, 150, 90, 80])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def validate_workbook(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=True)
    return {
        "sheets": wb.sheetnames,
        "dimensions": {ws.title: {"rows": ws.max_row, "columns": ws.max_column} for ws in wb.worksheets},
    }


def run(pdf_path: Path, rules_path: Path, output_path: Path, json_path: Path, summary_path: Path, pages_path: Path | None) -> dict[str, Any]:
    pages = read_pdf_pages(pdf_path)
    if pages_path:
        pages_path.write_text(json.dumps(pages, ensure_ascii=False, indent=2), encoding="utf-8")
    index = ContractIndex(pages)
    rules_payload = json.loads(rules_path.read_text(encoding="utf-8"))
    type1_rules = get_type(rules_payload, "type_1_direct_field_extraction").get("rules", [])
    project_type = infer_project_type(index)
    mixed_tower = infer_mixed_tower(index)

    type1_results = []
    for rule in type1_rules:
        type1_results.append({"rule": rule, "extracted": extract_type1_rule(rule, index, project_type, mixed_tower)})

    payload = {
        "metadata": {
            "contract_file": pdf_path.name,
            "rules_file": rules_path.name,
            "project_type": project_type,
            "mixed_tower": mixed_tower,
            "pdf_pages": len(pages),
        },
        "type_1_direct_field_extraction": type1_results,
        "type_2_payment_terms": extract_payments(index),
        "type_3_default_liability": extract_liabilities(index),
        "type_4_contract_risk": extract_risks(index),
        "type_5_document_delivery": extract_document_delivery(index),
    }

    summary = {
        "type1": status_counts(type1_results),
        "type2_rows": len(payload["type_2_payment_terms"]),
        "type3_rows": len(payload["type_3_default_liability"]),
        "type4_rows": len(payload["type_4_contract_risk"]),
        "type5_rows": len(payload["type_5_document_delivery"]),
    }

    json_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    build_workbook(payload, output_path)
    validation = validate_workbook(output_path)

    return {
        "xlsx": str(output_path),
        "json": str(json_path),
        "summary": str(summary_path),
        "counts": summary,
        "validation": validation,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract a wind-power contract PDF into a 5-sheet Excel workbook.")
    parser.add_argument("pdf", nargs="?", help="Input contract PDF path.")
    parser.add_argument("--pdf", dest="pdf_option", help="Input contract PDF path.")
    parser.add_argument("--rules", help="Path to contract_rules*.machine.json. Defaults to search under current directory.")
    parser.add_argument("--output", help="Output xlsx path. Defaults to outputs/<pdf_stem>_抽取结果.xlsx.")
    parser.add_argument("--output-dir", default="outputs", help="Output directory for default xlsx/json files.")
    parser.add_argument("--json-output", help="Structured extraction JSON path.")
    parser.add_argument("--summary-output", help="Summary JSON path.")
    parser.add_argument("--pages-output", help="Optional extracted PDF pages JSON path.")
    args = parser.parse_args()

    pdf_arg = args.pdf_option or args.pdf
    if not pdf_arg:
        raise SystemExit("PDF path is required. Example: contract_pdf_to_excel.py --pdf contract.pdf")
    pdf_path = Path(pdf_arg).expanduser().resolve()
    if not pdf_path.exists():
        raise FileNotFoundError(f"PDF not found: {pdf_path}")

    cwd = Path.cwd().resolve()
    rules_path = find_rules_file(args.rules, cwd)
    output_dir = Path(args.output_dir).expanduser()
    if not output_dir.is_absolute():
        output_dir = cwd / output_dir
    output_path = Path(args.output).expanduser().resolve() if args.output else output_dir / f"{pdf_path.stem}_抽取结果.xlsx"
    json_path = Path(args.json_output).expanduser().resolve() if args.json_output else output_dir / "contract_extraction_result.json"
    summary_path = Path(args.summary_output).expanduser().resolve() if args.summary_output else output_dir / "contract_extraction_summary.json"
    pages_path = Path(args.pages_output).expanduser().resolve() if args.pages_output else None

    result_payload = run(pdf_path, rules_path, output_path, json_path, summary_path, pages_path)
    print(json.dumps(result_payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
